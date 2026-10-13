"""Slides for the heart-failure readmission tutorial.

Reads figures and metrics.json written by the notebook. Does not query
BigQuery and does not open the cohort file.

    python make_hf_ml_slides.py build/hf_synthetic build/hf_synthetic/hf_readmission_ml_slides.pptx
"""
from __future__ import annotations

import json
import os
import re
import zipfile
from io import BytesIO

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# Same palette and type sizes as the other decks in this repo. Inlined so
# building slides does not import the BigQuery client.
W, H = Inches(13.333), Inches(7.5)
INK = RGBColor(0x1D, 0x24, 0x33)
MUTED = RGBColor(0x6B, 0x76, 0x84)
FAINT = RGBColor(0x9A, 0xA4, 0xB0)
RULE = RGBColor(0xDF, 0xE3, 0xE8)
ACCENT = RGBColor(0x2F, 0x6F, 0x9F)
GOLD = RGBColor(0xD9, 0x8F, 0x2B)
PANEL = RGBColor(0xF6, 0xF7, 0xF9)
SANS = "Calibri"
MARGIN = Inches(0.55)
CONTENT_W = W - 2 * MARGIN


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def textbox(slide, x, y, w, h, *, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def para(tf, text, *, size=16, bold=False, color=INK, first=False,
         space_before=0, space_after=0, line=1.15, align=PP_ALIGN.LEFT):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = line
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    run = p.add_run()
    run.text = text
    f = run.font
    f.name, f.size, f.bold = SANS, Pt(size), bold
    f.color.rgb = color
    return p


def rule(slide, x, y, w, *, color=ACCENT, h=Inches(0.045)):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    bar.shadow.inherit = False
    return bar


def panel(slide, x, y, w, h, *, fill=PANEL):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.shadow.inherit = False
    shape.adjustments[0] = 0.06
    return shape


def notes(slide, text):
    tf = slide.notes_slide.notes_text_frame
    tf.text = text.strip()


def kicker(slide, text):
    tf = textbox(slide, MARGIN, Inches(0.28), CONTENT_W, Inches(0.24))
    para(tf, text, size=12, bold=True, color=ACCENT, first=True)


def heading(slide, text):
    tf = textbox(slide, MARGIN, Inches(0.52), CONTENT_W, Inches(0.46))
    para(tf, text, size=26, bold=True, color=INK, first=True)


def footer(slide, text):
    tf = textbox(slide, MARGIN, H - Inches(0.36), CONTENT_W, Inches(0.22))
    para(tf, text, size=11, color=FAINT, first=True)


def add_fitted_picture(slide, path, x, y, box_w, box_h):
    from PIL import Image
    with Image.open(path) as im:
        pw, ph = im.size

    def as_in(v):
        return v.inches if hasattr(v, "inches") else v / 914400.0

    bw, bh = as_in(box_w), as_in(box_h)
    aspect = pw / float(ph)
    if aspect >= (bw / max(bh, 0.01)):
        w, h = Inches(bw), Inches(bw / aspect)
    else:
        h, w = Inches(bh), Inches(bh * aspect)
    return slide.shapes.add_picture(path, x, y, w, h)


def show_pic(slide, art, name, x, y, w, h):
    path = os.path.join(art, name)
    if os.path.exists(path):
        return add_fitted_picture(slide, path, x, y, w, h)
    tf = textbox(slide, x, y, w, Inches(0.4))
    para(tf, "Figure was not generated.", size=14, color=MUTED, first=True)
    return None


def load_metrics(art):
    with open(os.path.join(art, "metrics.json"), encoding="utf-8") as f:
        return json.load(f)


def pct(value):
    if value is None:
        return "—"
    return f"{100 * float(value):.1f}%"


def num(value, digits=2):
    if value is None:
        return "—"
    return f"{float(value):.{digits}f}"


def title_slide(prs, m):
    s = blank(prs)
    rule(s, MARGIN, Inches(1.85), Inches(1.2))
    tf = textbox(s, MARGIN, Inches(2.05), Inches(12), Inches(1.3))
    para(tf, "30-day readmission after heart failure", size=32, bold=True, first=True)
    where = "MIMIC-IV on BigQuery" if m["source"] == "bigquery" else "Synthetic rehearsal. Not a result."
    tf = textbox(s, MARGIN, Inches(3.45), Inches(11), Inches(1.2))
    para(
        tf,
        f"From a SQL cohort to four risk models. {m['n_stays']:,} stays, "
        f"{m['n_patients']:,} patients, {pct(m['readmit_rate'])} readmitted. {where}",
        size=18, color=MUTED, first=True,
    )
    footer(s, "All-cause readmission within 30 days of discharge")
    url = os.environ.get("HF_CODE_URL", "").strip() or "(set HF_CODE_URL and rebuild this deck)"
    rate_line = pct(m["readmit_rate"])
    notes(s, f"""
Code: {url}

Open with the question, not the model. {rate_line} of heart-failure discharges in this cohort are back within 30 days. The tutorial shows how to build that cohort without leaking the future, then how to compare a bedside score, a logistic regression, a gradient booster, and a small neural net.

Tell the room the source line on the slide. If it says synthetic, these numbers are a rehearsal and must not be quoted as MIMIC.

Walk them to the repository from the link above. The notebook comments are the lesson. The slides are the talk track.
""")


def objectives_slide(prs):
    s = blank(prs)
    kicker(s, "OUTCOMES")
    heading(s, "What you should be able to do")
    items = [
        ("1", "Define the cohort", "Heart-failure discharge, alive, not hospice, and a 30-day label that does not punish deaths."),
        ("2", "Keep the split honest", "Hold out patients, not stays, and show what a random split does to the score."),
        ("3", "Compare four models", "LACE, logistic regression, a gradient booster, and a neural net, on more than AUROC."),
        ("4", "Say what you would not ship", "Calibration, a staffing threshold, and where the score is weaker."),
    ]
    y = 1.35
    for num_label, title, body in items:
        panel(s, MARGIN, Inches(y), CONTENT_W, Inches(1.15))
        tf = textbox(s, MARGIN + Inches(0.25), Inches(y + 0.28), Inches(0.6), Inches(0.5))
        para(tf, num_label, size=20, bold=True, color=GOLD, first=True)
        tf = textbox(s, MARGIN + Inches(1.05), Inches(y + 0.18), Inches(10.5), Inches(0.4))
        para(tf, title, size=18, bold=True, first=True)
        tf = textbox(s, MARGIN + Inches(1.05), Inches(y + 0.58), Inches(10.5), Inches(0.4))
        para(tf, body, size=15, color=MUTED, first=True)
        y += 1.32
    notes(s, """
Read the four lines and stop. Do not preview the results yet.

Emphasize the second outcome. Most of the technical content is there to stop a leaked split and a model that looks good only because the AUROC went up.

The fourth outcome is the one to come back to at the end: a number you would not put in front of a case manager without the calibration plot and the subgroup check.
""")


def why_slide(prs, m):
    s = blank(prs)
    kicker(s, "QUESTION")
    heading(s, "Why this readmission")
    panel(s, MARGIN, Inches(1.4), Inches(3.6), Inches(2.3))
    tf = textbox(s, MARGIN + Inches(0.25), Inches(1.6), Inches(3.1), Inches(0.8))
    para(tf, pct(m["readmit_rate"]), size=36, bold=True, color=ACCENT, first=True)
    tf = textbox(s, MARGIN + Inches(0.25), Inches(2.5), Inches(3.1), Inches(0.8))
    para(tf, "of surviving discharges are back within 30 days", size=15, color=MUTED, first=True)
    points = [
        "Heart-failure readmission is a published quality measure and a staffing problem.",
        "The label here is all-cause. A planned return still counts. It is not the CMS measure.",
        "A useful score has to beat a bedside checklist, not just a coin flip.",
    ]
    tf = textbox(s, Inches(4.6), Inches(1.5), Inches(8), Inches(3.6))
    for i, line in enumerate(points):
        para(tf, line, size=18, first=(i == 0), space_after=14)
    notes(s, f"""
The rate on the slide is {pct(m['readmit_rate'])} of {m['n_stays']:,} discharges. Say the denominator out loud: these are people who left alive and were not sent to hospice.

Published heart-failure 30-day rates often land around 20 to 25 percent. If you are on the real extract, say whether you are in that range. If you are on the synthetic rehearsal, skip the comparison.

Stress the second bullet. Peers will ask if this is the Hospital Readmissions Reduction Program number. It is not. We did not remove planned readmissions, and the next admission can be for any reason.

The third bullet sets up LACE. We are not impressed by an AUROC of 0.60 until we see the checklist.
""")


def extends_slide(prs):
    s = blank(prs)
    kicker(s, "EXTENDS")
    heading(s, "How this extends the course tutorial")
    items = [
        ("Cohort", "Heart-failure discharges with a 30-day all-cause label, not a generic stay."),
        ("Honest split", "Hold out patients, not stays, and show what a random split does to AUROC."),
        ("Clinical baseline", "A LACE-style score sits next to the learned models."),
        ("Beyond AUROC", "Calibration, a decision curve, SHAP, and a sex and insurance audit."),
    ]
    y = 1.35
    for title, body in items:
        panel(s, MARGIN, Inches(y), CONTENT_W, Inches(1.15))
        tf = textbox(s, MARGIN + Inches(0.3), Inches(y + 0.18), Inches(11.5), Inches(0.36))
        para(tf, title, size=18, bold=True, first=True)
        tf = textbox(s, MARGIN + Inches(0.3), Inches(y + 0.58), Inches(11.5), Inches(0.4))
        para(tf, body, size=16, color=MUTED, first=True)
        y += 1.3
    notes(s, """
Say this once so peers know the assignment is not a copy of the lecture notebook.

The course example is often a readmission or length-of-stay walkthrough with a single model family. Here the cohort is heart-failure only, the split is by patient with a leakage demo, LACE is the bedside checklist, and the evaluation goes past AUROC into calibration, staffing thresholds, SHAP, and a fairness check.

Do not claim these methods are new. Claim that the teaching path is different and that a peer can rebuild it from the comments.
""")


def access_slide(prs):
    s = blank(prs)
    kicker(s, "DATA")
    heading(s, "Access, and what not to share")
    rows = [
        ("Credential", "MIMIC-IV on BigQuery, under the data use agreement you signed."),
        ("Billing", "Set GOOGLE_CLOUD_PROJECT to your project. The query is billed to you."),
        ("One download", "The notebook caches a parquet file. Do not commit it."),
        ("No row previews", "Slides and the PDF show counts and plots. Groups under 10 stays are dropped."),
    ]
    y = 1.4
    for title, body in rows:
        panel(s, MARGIN, Inches(y), CONTENT_W, Inches(1.15))
        tf = textbox(s, MARGIN + Inches(0.3), Inches(y + 0.18), Inches(11.5), Inches(0.36))
        para(tf, title, size=18, bold=True, first=True)
        tf = textbox(s, MARGIN + Inches(0.3), Inches(y + 0.58), Inches(11.5), Inches(0.4))
        para(tf, body, size=16, color=MUTED, first=True)
        y += 1.3
    notes(s, """
This slide is a rule, not a method. Read it before you open the notebook on a projector.

The cache and the executed notebook from a real run stay in the private folders. The repository link on the title slide is code only.

If a peer asks to see a row, the answer is no. The teaching artifacts are the query, the plots, and the metrics file, which has no identifiers.

labevents is scanned only for eight item ids. Say that removing the item-id filter is how a tutorial becomes an expensive, unnecessary scan.
""")


def pipeline_slide(prs):
    s = blank(prs)
    kicker(s, "PATH")
    heading(s, "The order the notebook runs")
    steps = [
        ("1", "Cohort", "Eligible discharges and a 30-day label."),
        ("2", "Split", "Patients, not stays. 70 / 15 / 15."),
        ("3", "Leakage check", "Random split versus patient split."),
        ("4", "Fit", "LACE, logistic, booster, neural net."),
        ("5", "Score once", "Test AUROC, precision-recall, Brier, calibration."),
        ("6", "Explain", "Threshold, decision curve, SHAP, subgroups."),
    ]
    positions = [(0, 0), (1, 0), (2, 0), (0, 1), (1, 1), (2, 1)]
    width, height, gap_x, gap_y = 3.9, 2.15, 0.2, 0.2
    origin_x, origin_y = 0.55, 1.4
    for (col, row), (num_label, title, body) in zip(positions, steps):
        x = origin_x + col * (width + gap_x)
        y = origin_y + row * (height + gap_y)
        panel(s, Inches(x), Inches(y), Inches(width), Inches(height))
        tf = textbox(s, Inches(x + 0.25), Inches(y + 0.25), Inches(width - 0.5), Inches(0.4))
        para(tf, num_label + "   " + title, size=18, bold=True, first=True)
        tf = textbox(s, Inches(x + 0.25), Inches(y + 0.85), Inches(width - 0.5), Inches(1.0))
        para(tf, body, size=15, color=MUTED, first=True)
    notes(s, """
Trace the six boxes with a finger. The rest of the talk is one box at a time.

Point at box 5 and say the test set is touched once. Box 3 does not use it. The booster's settings are chosen inside the training patients. The neural net is the only model that watches the validation loss.

Box 6 is where the talk earns the grade. Ranking is not a plan for who gets a phone call.
""")


def cohort_slide(prs, m):
    s = blank(prs)
    kicker(s, "COHORT")
    heading(s, "Who is in, and what the 1 means")
    left = [
        "Heart-failure diagnosis, ICD-10 I50 or ICD-9 428.",
        "Age 18 or older, left the hospital alive.",
        "Not discharged to hospice.",
        "Drop deaths within 30 days if they never returned.",
    ]
    right = [
        "LEAD() over every admission for that patient.",
        "1 if the next stay starts within 30 days.",
        "The next stay does not have to be heart failure.",
        "Final cohort: " + f"{m['n_stays']:,} stays.",
    ]
    panel(s, MARGIN, Inches(1.4), Inches(5.9), Inches(5.2))
    panel(s, Inches(6.85), Inches(1.4), Inches(5.9), Inches(5.2))
    tf = textbox(s, MARGIN + Inches(0.3), Inches(1.6), Inches(5.3), Inches(0.4))
    para(tf, "Index stay", size=18, bold=True, first=True)
    tf = textbox(s, MARGIN + Inches(0.3), Inches(2.2), Inches(5.3), Inches(4.0))
    for i, line in enumerate(left):
        para(tf, line, size=16, first=(i == 0), space_after=12)
    tf = textbox(s, Inches(7.15), Inches(1.6), Inches(5.3), Inches(0.4))
    para(tf, "Label", size=18, bold=True, first=True)
    tf = textbox(s, Inches(7.15), Inches(2.2), Inches(5.3), Inches(4.0))
    for i, line in enumerate(right):
        para(tf, line, size=16, first=(i == 0), space_after=12)
    notes(s, """
Walk the left column first. Each line is an exclusion, and the attrition chart in the notebook shows the counts. Hospice and inpatient death are not "failures to readmit." A death at day 10 with no return is also not a success, so those patients leave the denominator.

Then the right column. The easy mistake is to compute the next heart-failure admission only. That misses someone who comes back with pneumonia. LEAD() runs over all admissions, and only then do we keep the heart-failure index stays.

The window is 0 to 30 days from discharge, matching the SQL assignment's definition. Say that a same-day transfer can count, and that we did not try to drop planned readmissions.
""")


def features_slide(prs):
    s = blank(prs)
    kicker(s, "FEATURES")
    heading(s, "Only what you know at discharge")
    cards = [
        ("Stay", "Age, sex, insurance, admission type, where they came from, where they went, length of stay, ICU, weekend discharge."),
        ("History", "Admissions and emergency visits in the past year, days since the last discharge, comorbidity flags, a short Charlson score."),
        ("This stay", "First and last creatinine, BUN, sodium, potassium, hemoglobin, albumin, glucose, and NT-proBNP. Inpatient heart-failure meds."),
    ]
    x = 0.55
    for title, body in cards:
        panel(s, Inches(x), Inches(1.45), Inches(3.9), Inches(4.6))
        tf = textbox(s, Inches(x + 0.25), Inches(1.7), Inches(3.4), Inches(0.5))
        para(tf, title, size=20, bold=True, first=True)
        tf = textbox(s, Inches(x + 0.25), Inches(2.4), Inches(3.4), Inches(3.2))
        para(tf, body, size=15, color=MUTED, first=True)
        x += 4.15
    notes(s, """
The rule on the title is the whole slide. If a fact arrives after discharge, it is not a feature. That is why the labs are the first and last values charted during the stay, not a value from the next admission.

Days since the previous discharge is missing on a first stay. That missingness is a feature. The preprocessor adds an indicator instead of silently filling it.

Race is not in the model. Say that explicitly if someone asks. The audit later uses sex and insurance only.

Medication flags are inpatient orders. A home beta blocker that was never rewritten looks like "not ordered." That can make the medication coefficients hard to read. Mention it before someone over-interprets an odds ratio.
""")


def eda_slide(prs, art, m):
    s = blank(prs)
    kicker(s, "COHORT")
    heading(s, "The rate is not the same everywhere")
    show_pic(s, art, "disposition.png", MARGIN, Inches(1.3), Inches(7.4), Inches(5.4))
    panel(s, Inches(8.3), Inches(1.5), Inches(4.4), Inches(4.8))
    tf = textbox(s, Inches(8.55), Inches(1.75), Inches(4.0), Inches(4.3))
    para(tf, f"{m['n_stays']:,} stays", size=20, bold=True, first=True)
    para(tf, f"{m['n_patients']:,} patients", size=18, space_before=10)
    para(tf, f"{pct(m['readmit_rate'])} readmitted", size=18, space_before=8)
    n_miss = len(m.get("missing") or [])
    para(tf, f"{n_miss} numeric features have gaps", size=16, color=MUTED, space_before=16)
    para(tf, "Locations with fewer than 10 stays are omitted.", size=15, color=MUTED, space_before=12)
    notes(s, """
The chart is the raw readmission rate by discharge location, before any model. Use it to check that the label behaves. A location with a tiny denominator is omitted on purpose.

Do not invent a story the chart does not support. If skilled nursing and home are close, say they are close. The SQL assignment already found a smaller gap than people expect.

The missing-feature count is the reason the next modeling step keeps missingness indicators. Albumin and NT-proBNP are the usual holes. A missing BNP is often "not ordered," which is not the same as a random skip.
""")


def leakage_slide(prs, art, m):
    s = blank(prs)
    kicker(s, "SPLIT")
    heading(s, "Hold out patients, not stays")
    show_pic(s, art, "leakage.png", MARGIN, Inches(1.35), Inches(7.2), Inches(5.2))
    leak = m["leakage"]
    panel(s, Inches(8.1), Inches(1.7), Inches(4.6), Inches(4.4))
    tf = textbox(s, Inches(8.35), Inches(1.95), Inches(4.2), Inches(3.9))
    para(tf, "Random stays", size=16, color=MUTED, first=True)
    para(tf, num(leak["random_auroc"]), size=32, bold=True, space_before=4)
    para(tf, "Held-out patients", size=16, color=MUTED, space_before=16)
    para(tf, num(leak["grouped_auroc"]), size=32, bold=True, color=ACCENT, space_before=4)
    notes(s, f"""
Same booster, same features, training patients only. The test set is not involved.

Random stay split AUROC {num(leak['random_auroc'])}. Patient split AUROC {num(leak['grouped_auroc'])}.

If the random number is higher, the model is being graded on people it has already seen, because two admissions from one person fell on both sides. That gap is not a better algorithm.

The rest of the talk uses the patient split: about 70 percent train, 15 percent validation, 15 percent test, with no shared patient ids. The booster's own early-stopping split is turned off for the same reason. It would cut rows, not people.
""")


def prep_slide(prs):
    s = blank(prs)
    kicker(s, "PREPROCESS")
    heading(s, "One matrix for three of the models")
    steps = [
        ("Impute", "Median for numbers, most common level for categories. Fit on the training patients only."),
        ("Flag gaps", "A column that is 1 when the value was missing. First stays have no prior discharge."),
        ("Scale", "Zero mean, unit variance, so the logistic odds ratios are per standard deviation."),
        ("Encode", "One-hot categories, first level dropped, so each odds ratio has a reference."),
    ]
    y = 1.35
    for title, body in steps:
        panel(s, MARGIN, Inches(y), CONTENT_W, Inches(1.2))
        tf = textbox(s, MARGIN + Inches(0.3), Inches(y + 0.32), Inches(2.6), Inches(0.5))
        para(tf, title, size=18, bold=True, first=True)
        tf = textbox(s, Inches(3.6), Inches(y + 0.32), Inches(8.8), Inches(0.6))
        para(tf, body, size=16, color=MUTED, first=True)
        y += 1.35
    notes(s, """
Three models share this matrix: logistic regression, the booster, and the network. LACE does not. It uses the raw length of stay, the admission type, the Charlson score, and the six-month emergency count.

Fitting the imputer on everybody, including the test patients, is a quieter leak than the split we just showed. The notebook fits it inside each training fold, then once on train plus validation for the final logistic and booster.

The booster could have taken the missing values directly. We still impute, so the comparison is about the model and not about a different matrix.

Scaling is for the logistic model and the network. Mention that an odds ratio is not "per year of age." It is per standard deviation, or versus the dropped category.
""")


def models_slide(prs, m):
    s = blank(prs)
    kicker(s, "MODELS")
    heading(s, "A checklist, a line, and a booster")
    params = m.get("gbm_params") or {}
    cards = [
        ("LACE", "Points for stay length, an emergency admission, Charlson band, and recent emergency visits. One logistic turns points into a probability."),
        ("Logistic", "L2 penalty, C = 1, no class weight. Odds ratios come from this fit. About a quarter of stays are events, so we do not reweight."),
        ("Booster", f"Grouped cross-validation on the training patients. Chosen depth {params.get('max_depth', '—')}, learning rate {params.get('learning_rate', '—')}. Refit on train plus validation."),
    ]
    x = 0.55
    for title, body in cards:
        panel(s, Inches(x), Inches(1.45), Inches(3.9), Inches(4.8))
        tf = textbox(s, Inches(x + 0.25), Inches(1.7), Inches(3.4), Inches(0.5))
        para(tf, title, size=20, bold=True, first=True)
        tf = textbox(s, Inches(x + 0.25), Inches(2.45), Inches(3.4), Inches(3.4))
        para(tf, body, size=15, color=MUTED, first=True)
        x += 4.15
    notes(s, f"""
LACE is the baseline a clinician already understands. Be honest that the Charlson piece is shortened. The SQL file lists the weights. Do not call the total the published LACE index.

The logistic model is the interpretable full model. We left class_weight off so the probabilities stay on the calibration plot. With a 25 percent outcome, balance is not the problem.

The booster's settings were picked by 3-fold grouped cross-validation. Mean validation AUROC was {num(m.get('gbm_cv_auroc'))}. Early stopping inside the booster is off, because that split is by row.

All three are refit on train plus validation before the test score. The network, next slide, is not, because it used validation to decide when to stop.
""")


def mlp_slide(prs, m):
    s = blank(prs)
    kicker(s, "NEURAL NET")
    heading(s, "A small network, stopped on validation loss")
    lines = [
        ("Input", "The same imputed and scaled matrix the logistic model sees."),
        ("Shape", "Linear 64, ReLU, dropout 0.2, linear 32, ReLU, one logit."),
        ("Loss", "Binary cross-entropy on logits. No positive-class weight."),
        ("Stop", f"Lowest validation loss, then four quiet epochs. Stopped at epoch {m.get('mlp_best_epoch', '—')}."),
    ]
    y = 1.35
    for title, body in lines:
        panel(s, MARGIN, Inches(y), CONTENT_W, Inches(1.2))
        tf = textbox(s, MARGIN + Inches(0.3), Inches(y + 0.32), Inches(2.2), Inches(0.5))
        para(tf, title, size=18, bold=True, first=True)
        tf = textbox(s, Inches(3.3), Inches(y + 0.32), Inches(9.2), Inches(0.6))
        para(tf, body, size=16, color=MUTED, first=True)
        y += 1.35
    notes(s, f"""
Show the training loop in the notebook while this slide is up. The points to narrate: the preprocessor is fit on train only, dropout is on during training and off when we score, and we restore the weights from the best validation epoch rather than keeping the last one.

Best epoch {m.get('mlp_best_epoch', '—')}, validation loss {num(m.get('mlp_val_loss'), 3)}.

We did not refit this model on train plus validation. Doing that would throw away the signal we used to stop. The other three models do not have that constraint.

If someone asks why the network is not bigger: the table is a few dozen columns, not an image. A deeper net memorizes the training patients. The exercise at the end is to add a positive-class weight and watch the calibration plot move.
""")


def compare_slide(prs, art, m):
    s = blank(prs)
    kicker(s, "TEST SET")
    heading(s, "Rank, then check the probability")
    show_pic(s, art, "roc.png", MARGIN, Inches(1.25), Inches(7.0), Inches(5.5))
    panel(s, Inches(7.8), Inches(1.4), Inches(4.9), Inches(5.3))
    tf = textbox(s, Inches(8.05), Inches(1.55), Inches(4.5), Inches(0.4))
    para(tf, "AUROC  (95% bootstrap)", size=14, bold=True, color=MUTED, first=True)
    y = 2.1
    for name in ("LACE", "Logistic", "GBM", "MLP"):
        row = m["models"][name]
        tf = textbox(s, Inches(8.05), Inches(y), Inches(4.5), Inches(0.9))
        para(tf, name, size=14, color=MUTED, first=True)
        para(tf, f"{num(row['auroc'])}   {num(row['auroc_lo'])}–{num(row['auroc_hi'])}", size=20, bold=True, space_before=2)
        y += 1.05
    notes(s, """
Read the four AUROCs and their intervals. Overlapping intervals means you cannot crown a winner from this one test set.

AUROC is a ranking. It does not say that a predicted 0.4 is a real 40 percent. That is the next slide.

Average precision is in the notebook next to the ROC, with the prevalence drawn as a horizontal line. Mention it if the room only stares at the ROC: a rare outcome can have a pretty ROC and a weak precision-recall curve. This outcome is common enough that both should move together.

These scores are the single look at the test patients. Do not go back and retune.
""")


def calibration_slide(prs, art, m):
    s = blank(prs)
    kicker(s, "CALIBRATION")
    heading(s, "Does 30% actually happen 30% of the time?")
    show_pic(s, art, "calibration.png", MARGIN, Inches(1.25), Inches(7.2), Inches(5.5))
    panel(s, Inches(8.0), Inches(1.5), Inches(4.7), Inches(5.0))
    tf = textbox(s, Inches(8.25), Inches(1.7), Inches(4.3), Inches(0.4))
    para(tf, "Brier score, lower is better", size=14, bold=True, color=MUTED, first=True)
    y = 2.3
    for name in ("LACE", "Logistic", "GBM", "MLP"):
        row = m["models"][name]
        tf = textbox(s, Inches(8.25), Inches(y), Inches(4.3), Inches(0.85))
        para(tf, f"{name}    {num(row['brier'], 3)}", size=18, bold=True, first=True)
        para(tf, f"Average precision {num(row['auprc'])}", size=14, color=MUTED, space_before=2)
        y += 0.95
    notes(s, """
The diagonal is perfect calibration. A curve above the line means the model is too humble in that bin: events happen more often than it said. Below the line, it is overconfident.

Brier score punishes confident mistakes. A model can win on AUROC and lose on Brier. If that happens here, say so, and do not recommend its raw probabilities for a threshold.

We kept class weights off so this plot would be readable. Adding a positive-class weight pulls probabilities toward the middle or the top and usually makes this picture worse. That is an exercise in the notebook, not a bug.
""")


def decision_slide(prs, art, m):
    s = blank(prs)
    kicker(s, "ACTION")
    heading(s, "Who would a case manager call?")
    show_pic(s, art, "dca.png", MARGIN, Inches(1.25), Inches(7.4), Inches(5.5))
    op = m["operating_point"]
    panel(s, Inches(8.15), Inches(1.55), Inches(4.55), Inches(4.7))
    tf = textbox(s, Inches(8.4), Inches(1.8), Inches(4.15), Inches(4.2))
    para(tf, "Top fifth, booster", size=16, color=MUTED, first=True)
    para(tf, pct(op["precision"]), size=32, bold=True, color=GOLD, space_before=6)
    para(tf, "of calls are true readmissions", size=15, color=MUTED, space_before=2)
    para(tf, pct(op["recall"]), size=32, bold=True, space_before=14)
    para(tf, "of all readmissions are caught", size=15, color=MUTED, space_before=2)
    notes(s, f"""
The operating point is concrete. Flag the highest-risk 20 percent under the booster. Precision {pct(op['precision'])}: of the people you call, this fraction come back. Recall {pct(op['recall'])}: of everyone who comes back, this fraction was on your list. Prevalence is {pct(op['prevalence'])}. If precision is near the prevalence, the list is not enriching risk.

The decision curve is the more general version. The horizontal axis is the threshold probability at which you would bother to intervene. Net benefit is true positives minus a penalty for false positives. The model is useful only where its line beats both zero and "flag everyone."

LACE is on the same plot so you can see whether the booster changes the decision, not just the AUROC.
""")


def shap_slide(prs, art, m):
    s = blank(prs)
    kicker(s, "EXPLANATION")
    heading(s, "What the booster is using")
    show_pic(s, art, "shap.png", MARGIN, Inches(1.25), Inches(7.6), Inches(5.5))
    panel(s, Inches(8.35), Inches(1.5), Inches(4.4), Inches(5.0))
    tf = textbox(s, Inches(8.55), Inches(1.7), Inches(4.0), Inches(0.4))
    para(tf, "Largest logistic odds ratios", size=14, bold=True, color=MUTED, first=True)
    tf = textbox(s, Inches(8.55), Inches(2.25), Inches(4.0), Inches(3.9))
    first = True
    for row in (m.get("odds_ratios") or [])[:5]:
        para(
            tf,
            row["feature"].replace("_", " "),
            size=14, bold=True, first=first, space_before=0 if first else 8,
        )
        first = False
        para(tf, num(row["odds_ratio"]), size=16, color=ACCENT, space_before=0)
    notes(s, """
The bars are mean absolute SHAP on a sample of test stays. They say how much a feature moves the booster's score. They do not say whether the feature pushes risk up or down. For direction, use the logistic odds ratios on the right, and remember they are per standard deviation or versus a reference category.

Do not narrate a medication coefficient as a treatment effect. Inpatient orders are not home regimens, and sicker patients get more diuretics.

The network has no trees. The notebook reports permutation importance for it: shuffle a column and watch the AUROC drop. If the top features disagree with the booster, say they disagree. The two models are allowed to use the table differently.
""")


def subgroup_slide(prs, art, m):
    s = blank(prs)
    kicker(s, "AUDIT")
    heading(s, "Same score, by sex and insurance")
    show_pic(s, art, "subgroups.png", MARGIN, Inches(1.3), Inches(8.2), Inches(5.4))
    panel(s, Inches(9.0), Inches(1.7), Inches(3.7), Inches(4.4))
    tf = textbox(s, Inches(9.2), Inches(1.95), Inches(3.35), Inches(3.9))
    para(tf, "Pre-specified", size=16, bold=True, first=True)
    para(tf, "The audit uses the booster, chosen before looking at which model won.", size=15, color=MUTED, space_before=8)
    para(tf, "Levels under 10 stays, or under 10 events, are omitted.", size=15, color=MUTED, space_before=10)
    notes(s, """
The gold line on the figure, if you can see it, is the overall test AUROC. A group sitting well below that line is where the score is less informative.

We required 10 stays, 10 readmissions, and 10 non-readmissions. Smaller cells are both unstable and too specific to show.

This is not a fairness certificate. It is a check that one large group is not carrying the whole AUROC. Race was not a feature and is not sliced here.

Do not propose a different cutoff by insurance from this one chart. That is a clinical and policy decision the model does not make.
""")


def limits_slide(prs):
    s = blank(prs)
    kicker(s, "LIMITS")
    heading(s, "What you should not claim")
    items = [
        "All-cause label. Planned returns count. Not a CMS measure.",
        "One health system, shifted dates, no external validation.",
        "ICD text match, not chart review. Inpatient orders miss home meds.",
        "Shortened Charlson score inside the LACE baseline.",
        "Not a decision tool. A high score is a reason to look.",
    ]
    y = 1.35
    for i, line in enumerate(items, start=1):
        panel(s, MARGIN, Inches(y), CONTENT_W, Inches(0.95))
        tf = textbox(s, MARGIN + Inches(0.25), Inches(y + 0.25), Inches(0.6), Inches(0.45))
        para(tf, str(i), size=18, bold=True, color=GOLD, first=True)
        tf = textbox(s, MARGIN + Inches(1.0), Inches(y + 0.25), Inches(10.8), Inches(0.5))
        para(tf, line, size=16, first=True)
        y += 1.05
    notes(s, """
Read these as the claims you refuse, not as a footnote. The one peers skip is the first: they will want to call it the CMS readmission rate. It is not.

The Charlson line protects you if you show LACE points. The weights in the SQL are a subset, and metastatic cancer replaces the solid-tumor weight rather than stacking on it.

Close by repeating that the score is not an order. The decision curve tells you where flagging helps relative to flagging everyone. It does not tell you the phone script.
""")


def replicate_slide(prs):
    s = blank(prs)
    kicker(s, "REPLICATE")
    heading(s, "Run it in this order")
    lines = [
        "Create a Python 3.12 environment and install requirements-ml.txt.",
        "Dry run without MIMIC: run_hf_ml.py --source synthetic.",
        "Set GOOGLE_CLOUD_PROJECT to your billing project and sign in.",
        "Run the tutorial script with the real-data source.",
        "Put the anonymous code link in HF_CODE_URL and rebuild the title notes.",
    ]
    y = 1.35
    for i, line in enumerate(lines, start=1):
        tf = textbox(s, MARGIN, Inches(y), CONTENT_W, Inches(0.85))
        para(tf, f"{i}    {line}", size=18, first=True)
        y += 0.95
    notes(s, """
Point at the repository from the title slide rather than reading install commands off this slide.

Step 2 is for peers without PhysioNet credentials. Synthetic numbers must not be quoted as MIMIC.

The real-data command is in the README. Do not paste a billing-project id into the deck or the code.

After a real run, the notebook, the PDF, and this deck are written under the private build folder. Copy the PowerPoint and the PDF out to submit them. Leave the parquet cache where it is.

If a cell fails, rerun the synthetic source to see the traceback. Do not paste a real error that includes a row.
""")


def try_yourself_slide(prs):
    s = blank(prs)
    kicker(s, "PRACTICE")
    heading(s, "Try it yourself")
    items = [
        ("1", "Swap the booster", "Replace HistGradientBoostingClassifier with a RandomForestClassifier and compare AUROC and calibration."),
        ("2", "Move the threshold", "Change the call-list cutoff on the decision curve and say who gets contacted."),
        ("3", "Add one lab", "Pull one more lab item id into the cohort query and refit the logistic model."),
    ]
    y = 1.4
    for num_label, title, body in items:
        panel(s, MARGIN, Inches(y), CONTENT_W, Inches(1.4))
        tf = textbox(s, MARGIN + Inches(0.25), Inches(y + 0.28), Inches(0.6), Inches(0.5))
        para(tf, num_label, size=22, bold=True, color=GOLD, first=True)
        tf = textbox(s, MARGIN + Inches(1.05), Inches(y + 0.22), Inches(10.5), Inches(0.4))
        para(tf, title, size=18, bold=True, first=True)
        tf = textbox(s, MARGIN + Inches(1.05), Inches(y + 0.7), Inches(10.5), Inches(0.5))
        para(tf, body, size=16, color=MUTED, first=True)
        y += 1.55
    notes(s, """
These three exercises are also at the end of the notebook. Pause here if the room has time to open the code.

The first exercise keeps the pipeline and only changes the model class. The second forces a staffing conversation instead of chasing AUROC. The third is the smallest feature change that still touches the SQL.

If time is short, assign one exercise as homework and leave the slide up.
""")


def references_slide(prs):
    s = blank(prs)
    kicker(s, "SOURCES")
    heading(s, "References")
    refs = [
        "van Walraven et al. LACE index. CMAJ, 2010.",
        "CMS Hospital Readmissions Reduction Program. The label here is not that measure.",
        "Johnson et al. MIMIC-IV. PhysioNet.",
        "MIT-LCP mimic-code repository, for cohort conventions.",
        "Vickers and Elkin. Decision curve analysis. Medical Decision Making, 2006.",
        "Lundberg and Lee. SHAP. NeurIPS, 2017.",
    ]
    tf = textbox(s, MARGIN, Inches(1.45), Inches(12), Inches(5.2))
    for i, line in enumerate(refs):
        para(tf, line, size=18, first=(i == 0), space_after=12)
    notes(s, """
Leave this slide up for questions.

If someone asks for the LACE paper, it is the 2010 CMAJ derivation. Remind them our Charlson inputs are shortened, so the points are a baseline in that style.

If someone asks for the data paper, it is MIMIC-IV on PhysioNet. The mimic-code repository is the usual place for the diagnosis regexes and the age formula: anchor age plus the year offset. Dates are shifted into the future on purpose.

Decision curves are Vickers and Elkin 2006. SHAP is Lundberg and Lee 2017. You do not need to derive either. You need to say what the picture is for: net benefit at a threshold, and which features move the booster.
""")


def scrub_author(path):
    """Drop creator names python-pptx writes from the OS user."""
    buf = BytesIO()
    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(buf, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/core.xml":
                text = data.decode("utf-8")
                text = re.sub(r"<dc:creator>.*?</dc:creator>", "<dc:creator></dc:creator>", text)
                text = re.sub(
                    r"<cp:lastModifiedBy>.*?</cp:lastModifiedBy>",
                    "<cp:lastModifiedBy></cp:lastModifiedBy>",
                    text,
                )
                data = text.encode("utf-8")
            zout.writestr(item, data)
    with open(path, "wb") as f:
        f.write(buf.getvalue())


def build(art, out):
    m = load_metrics(art)
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    prs.core_properties.author = ""
    prs.core_properties.last_modified_by = ""
    prs.core_properties.title = "Heart-failure readmission tutorial"
    prs.core_properties.subject = ""
    prs.core_properties.comments = ""
    title_slide(prs, m)
    objectives_slide(prs)
    why_slide(prs, m)
    extends_slide(prs)
    access_slide(prs)
    pipeline_slide(prs)
    cohort_slide(prs, m)
    features_slide(prs)
    eda_slide(prs, art, m)
    leakage_slide(prs, art, m)
    prep_slide(prs)
    models_slide(prs, m)
    mlp_slide(prs, m)
    compare_slide(prs, art, m)
    calibration_slide(prs, art, m)
    decision_slide(prs, art, m)
    shap_slide(prs, art, m)
    subgroup_slide(prs, art, m)
    limits_slide(prs)
    replicate_slide(prs)
    try_yourself_slide(prs)
    references_slide(prs)
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    prs.save(out)
    scrub_author(out)
    print("wrote", out, "slides", len(prs.slides))


if __name__ == "__main__":
    import sys
    art = sys.argv[1]
    out = sys.argv[2]
    build(art, out)
